from docx import Document


def extract_text_from_docx(file_path):

    document = Document(file_path)

    text_parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells
            )
            text_parts.append(row_text)

    return "\n".join(text_parts)